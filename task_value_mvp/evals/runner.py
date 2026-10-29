from __future__ import annotations

import base64
import hashlib
import html
import ipaddress
import json
import os
import re
import shutil
import socket
import subprocess
import tempfile
import time
import urllib.error
import urllib.parse
import urllib.request
from collections import Counter
from dataclasses import asdict, dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import yaml

from task_value_mvp.config import ARTIFACT_DIR
from task_value_mvp.catalog import MODEL_CATALOG

from .graders import grade_bfcl, grade_gaia, grade_spreadsheet
from .manifest import DEFAULT_CACHE_DIR, load_built_tasks
from .schema import EvalResult, EvalTask, GradeResult, ModelConfig


@dataclass(frozen=True)
class RunConfig:
    max_cost_per_task_usd: float = 1.0
    max_output_tokens: int = 2048
    max_tool_calls: int = 12
    max_tool_output_chars: int = 20_000
    timeout_seconds: int = 90


@dataclass
class ApiResponse:
    text: str
    tool_calls: list[dict[str, Any]]
    raw_message: dict[str, Any]
    input_tokens: int
    output_tokens: int
    cached_input_tokens: int
    latency_ms: float


def load_models_config(path: Path) -> tuple[RunConfig, list[ModelConfig]]:
    value = yaml.safe_load(path.read_text(encoding="utf-8")) or {}
    run = RunConfig(**(value.get("run") or {}))
    models = [ModelConfig(**item) for item in value.get("models", [])]
    if len(models) != 3:
        raise ValueError("Exactly three eval models are required")
    if {model.role for model in models} != {"low", "middle", "premium"}:
        raise ValueError("Eval models must have unique low, middle, and premium roles")
    if len({model.alias for model in models}) != 3:
        raise ValueError("Eval model aliases must be unique")
    catalog_ids = {item.model_id for item in MODEL_CATALOG}
    unknown = [model.catalog_model_id for model in models if model.catalog_model_id not in catalog_ids]
    if unknown:
        raise ValueError(f"Eval models have unknown catalog IDs: {', '.join(unknown)}")
    if len({model.catalog_model_id for model in models}) != 3:
        raise ValueError("Eval models must map to three distinct catalog IDs")
    missing = [model.api_key_env for model in models if not os.environ.get(model.api_key_env)]
    if missing:
        raise ValueError(f"Missing model API credentials in environment: {', '.join(missing)}")
    return run, models


class OpenAICompatibleClient:
    def __init__(self, model: ModelConfig, timeout_seconds: int) -> None:
        self.model = model
        self.timeout_seconds = timeout_seconds

    def complete(
        self,
        messages: list[dict[str, Any]],
        tools: list[dict[str, Any]] | None,
        max_output_tokens: int,
    ) -> ApiResponse:
        payload: dict[str, Any] = {
            "model": self.model.api_model,
            "messages": messages,
            "temperature": 0,
            "seed": 42,
            "max_tokens": max_output_tokens,
        }
        if tools:
            payload["tools"] = tools
            payload["tool_choice"] = "auto"
        endpoint = self.model.base_url.rstrip("/") + "/chat/completions"
        request = urllib.request.Request(
            endpoint,
            data=json.dumps(payload).encode("utf-8"),
            headers={
                "Authorization": f"Bearer {os.environ[self.model.api_key_env]}",
                "Content-Type": "application/json",
                "User-Agent": "task-router-evals/1",
            },
            method="POST",
        )
        started = time.perf_counter()
        last_error: Exception | None = None
        for attempt in range(3):
            try:
                with urllib.request.urlopen(request, timeout=self.timeout_seconds) as response:
                    value = json.loads(response.read().decode("utf-8"))
                break
            except urllib.error.HTTPError as exc:
                last_error = exc
                if exc.code != 429 and exc.code < 500:
                    raise RuntimeError(f"Model API returned HTTP {exc.code}") from exc
            except (urllib.error.URLError, TimeoutError) as exc:
                last_error = exc
            if attempt == 2:
                raise RuntimeError(f"Model API failed after retries: {last_error}") from last_error
            time.sleep(0.25 * (2**attempt))
        latency = (time.perf_counter() - started) * 1000
        try:
            message = value["choices"][0]["message"]
        except (KeyError, IndexError, TypeError) as exc:
            raise RuntimeError("Model API response has no assistant message") from exc
        usage = value.get("usage") or {}
        details = usage.get("prompt_tokens_details") or {}
        calls = []
        for call in message.get("tool_calls") or []:
            function = call.get("function") or {}
            calls.append(
                {
                    "id": call.get("id"),
                    "name": function.get("name", ""),
                    "arguments": function.get("arguments", "{}"),
                }
            )
        return ApiResponse(
            text=message.get("content") or "",
            tool_calls=calls,
            raw_message=message,
            input_tokens=int(usage.get("prompt_tokens") or usage.get("input_tokens") or 0),
            output_tokens=int(usage.get("completion_tokens") or usage.get("output_tokens") or 0),
            cached_input_tokens=int(details.get("cached_tokens") or usage.get("cached_input_tokens") or 0),
            latency_ms=latency,
        )


def _openai_tools(functions: list[dict[str, Any]]) -> list[dict[str, Any]]:
    tools = []
    for function in functions:
        value = json.loads(json.dumps(function))
        if value.get("parameters", {}).get("type") == "dict":
            value["parameters"]["type"] = "object"
        value["name"] = value["name"].replace(".", "_")
        tools.append({"type": "function", "function": value})
    return tools


def _workbook_summary(path: Path, max_rows: int = 18, max_columns: int = 14) -> str:
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=False, read_only=True)
    sections = []
    for sheet in workbook.worksheets[:8]:
        rows = []
        for row in sheet.iter_rows(
            min_row=1, max_row=min(sheet.max_row, max_rows), max_col=min(sheet.max_column, max_columns), values_only=True
        ):
            rows.append(list(row))
        sections.append(
            f"Sheet {sheet.title!r} ({sheet.max_row}x{sheet.max_column}) sample:\n"
            + json.dumps(rows, default=str)[:8000]
        )
    return "\n\n".join(sections)


def _asset_content(task: EvalTask) -> list[dict[str, Any]]:
    parts: list[dict[str, Any]] = [{"type": "text", "text": task.prompt}]
    for value in task.assets:
        path = Path(value)
        suffix = path.suffix.lower()
        if suffix in {".png", ".jpg", ".jpeg"}:
            mime = "image/png" if suffix == ".png" else "image/jpeg"
            encoded = base64.b64encode(path.read_bytes()).decode("ascii")
            parts.append({"type": "image_url", "image_url": {"url": f"data:{mime};base64,{encoded}"}})
        elif suffix == ".xlsx":
            parts.append({"type": "text", "text": f"Attached workbook {path.name}:\n{_workbook_summary(path)}"})
        else:
            parts.append(
                {"type": "text", "text": f"Attached file {path.name}:\n{path.read_text(errors='replace')[:20000]}"}
            )
    parts.append({"type": "text", "text": "Return only the final answer, without explanation."})
    return parts


SPREADSHEET_TOOL = {
    "type": "function",
    "function": {
        "name": "spreadsheet_action",
        "description": "Inspect or modify the working Excel workbook. Call finish when the workbook is complete.",
        "parameters": {
            "type": "object",
            "properties": {
                "operation": {
                    "type": "string",
                    "enum": ["list_sheets", "read_range", "write_cells", "copy_range", "insert_rows", "delete_rows", "sort_range", "finish"],
                },
                "sheet": {"type": "string"},
                "range": {"type": "string", "description": "Excel range such as A1:D20"},
                "destination": {"type": "string"},
                "cells": {
                    "type": "array",
                    "items": {
                        "type": "object",
                        "properties": {"cell": {"type": "string"}, "value": {}},
                        "required": ["cell", "value"],
                    },
                },
                "row": {"type": "integer"},
                "amount": {"type": "integer"},
                "key_columns": {"type": "array", "items": {"type": "integer"}},
                "reverse": {"type": "boolean"},
            },
            "required": ["operation"],
        },
    },
}

GAIA_TOOLS = [
    {
        "type": "function",
        "function": {
            "name": "web_search",
            "description": "Search the public web for sources relevant to the question.",
            "parameters": {
                "type": "object",
                "properties": {"query": {"type": "string"}},
                "required": ["query"],
            },
        },
    },
    {
        "type": "function",
        "function": {
            "name": "web_fetch",
            "description": "Fetch readable text from one public HTTP or HTTPS URL returned by web_search.",
            "parameters": {
                "type": "object",
                "properties": {"url": {"type": "string"}},
                "required": ["url"],
            },
        },
    },
]


def _validate_public_url(value: str) -> str:
    parsed = urllib.parse.urlparse(value)
    if parsed.scheme not in {"http", "https"} or not parsed.hostname or parsed.username or parsed.password:
        raise ValueError("Only public HTTP and HTTPS URLs are allowed")
    if parsed.port not in {None, 80, 443}:
        raise ValueError("Non-standard URL ports are not allowed")
    addresses = socket.getaddrinfo(parsed.hostname, parsed.port or (443 if parsed.scheme == "https" else 80))
    if not addresses or any(not ipaddress.ip_address(item[4][0]).is_global for item in addresses):
        raise ValueError("Private or non-routable URL targets are not allowed")
    return value


class _PublicRedirectHandler(urllib.request.HTTPRedirectHandler):
    def redirect_request(self, request, file_pointer, code, message, headers, new_url):
        _validate_public_url(new_url)
        return super().redirect_request(request, file_pointer, code, message, headers, new_url)


def _web_fetch(url: str, timeout_seconds: int, limit: int) -> str:
    _validate_public_url(url)
    request = urllib.request.Request(url, headers={"User-Agent": "task-router-evals/1"})
    opener = urllib.request.build_opener(_PublicRedirectHandler())
    with opener.open(request, timeout=timeout_seconds) as response:
        payload = response.read(1_000_001)
    if len(payload) > 1_000_000:
        raise ValueError("Fetched document exceeds 1 MB")
    text = payload.decode("utf-8", errors="replace")
    text = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", text)
    text = re.sub(r"(?s)<[^>]+>", " ", text)
    return re.sub(r"\s+", " ", html.unescape(text)).strip()[:limit]


def _web_search(query: str, timeout_seconds: int, limit: int) -> str:
    if not query.strip():
        raise ValueError("Search query is empty")
    url = "https://html.duckduckgo.com/html/?" + urllib.parse.urlencode({"q": query[:500]})
    request = urllib.request.Request(url, headers={"User-Agent": "task-router-evals/1"})
    with urllib.request.urlopen(request, timeout=timeout_seconds) as response:
        payload = response.read(500_001)
    if len(payload) > 500_000:
        raise ValueError("Search response exceeds 500 KB")
    source = payload.decode("utf-8", errors="replace")
    rows = []
    pattern = re.compile(r'<a[^>]+class="[^"]*result__a[^"]*"[^>]+href="([^"]+)"[^>]*>(.*?)</a>', re.I | re.S)
    for href, title in pattern.findall(source)[:8]:
        parsed = urllib.parse.urlparse(html.unescape(href))
        target = urllib.parse.parse_qs(parsed.query).get("uddg", [html.unescape(href)])[0]
        clean_title = re.sub(r"<[^>]+>", "", html.unescape(title))
        rows.append(f"{clean_title.strip()}\n{target}")
    return ("\n\n".join(rows) or "No search results found.")[:limit]


def _remaining_timeout(client: OpenAICompatibleClient, deadline: float) -> int:
    remaining = deadline - time.monotonic()
    if remaining <= 0:
        raise TimeoutError("Task time limit exceeded")
    client.timeout_seconds = max(1, min(client.timeout_seconds, int(remaining) + 1))
    return client.timeout_seconds


def _spreadsheet_action(path: Path, arguments: dict[str, Any]) -> dict[str, Any]:
    import openpyxl
    from copy import copy
    from openpyxl.utils.cell import range_boundaries

    workbook = openpyxl.load_workbook(path)
    operation = arguments.get("operation")
    if operation == "list_sheets":
        return {"sheets": [{"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column} for sheet in workbook]}
    if operation == "finish":
        return {"finished": True}
    sheet_name = arguments.get("sheet")
    if sheet_name not in workbook.sheetnames:
        return {"error": f"unknown sheet: {sheet_name}"}
    sheet = workbook[sheet_name]
    if operation == "read_range":
        min_col, min_row, max_col, max_row = range_boundaries(arguments["range"])
        if (max_col - min_col + 1) * (max_row - min_row + 1) > 500:
            return {"error": "read range exceeds 500 cells"}
        values = [
            [sheet.cell(row=row, column=column).value for column in range(min_col, max_col + 1)]
            for row in range(min_row, max_row + 1)
        ]
        return {"range": arguments["range"], "values": values}
    if operation == "write_cells":
        cells = arguments.get("cells") or []
        if len(cells) > 500:
            return {"error": "write exceeds 500 cells"}
        for item in cells:
            sheet[item["cell"]] = item.get("value")
    elif operation == "copy_range":
        min_col, min_row, max_col, max_row = range_boundaries(arguments["range"])
        destination = sheet[arguments["destination"]]
        for row_offset, row in enumerate(range(min_row, max_row + 1)):
            for col_offset, column in enumerate(range(min_col, max_col + 1)):
                source = sheet.cell(row=row, column=column)
                target = sheet.cell(row=destination.row + row_offset, column=destination.column + col_offset)
                target.value = source.value
                if source.has_style:
                    target._style = copy(source._style)
    elif operation == "insert_rows":
        sheet.insert_rows(int(arguments["row"]), amount=max(1, int(arguments.get("amount", 1))))
    elif operation == "delete_rows":
        sheet.delete_rows(int(arguments["row"]), amount=max(1, int(arguments.get("amount", 1))))
    elif operation == "sort_range":
        min_col, min_row, max_col, max_row = range_boundaries(arguments["range"])
        rows = [[sheet.cell(row=row, column=column).value for column in range(min_col, max_col + 1)] for row in range(min_row, max_row + 1)]
        keys = [max(1, int(value)) - 1 for value in arguments.get("key_columns") or [1]]
        rows.sort(key=lambda row: tuple((row[index] is None, str(row[index])) for index in keys), reverse=bool(arguments.get("reverse")))
        for row_offset, values in enumerate(rows):
            for col_offset, value in enumerate(values):
                sheet.cell(row=min_row + row_offset, column=min_col + col_offset).value = value
    else:
        return {"error": f"unsupported operation: {operation}"}
    workbook.save(path)
    return {"ok": True, "operation": operation}


def _recalculate_workbook(path: Path, deadline: float) -> bool:
    binary = shutil.which("libreoffice") or shutil.which("soffice")
    mac_binary = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    if not binary and mac_binary.exists():
        binary = str(mac_binary)
    if not binary:
        return False
    remaining = deadline - time.monotonic()
    if remaining <= 0:
        raise TimeoutError("Task time limit exceeded before workbook recalculation")
    with tempfile.TemporaryDirectory(prefix="task-router-recalc-") as temporary:
        completed = subprocess.run(
            [binary, "--headless", "--calc", "--convert-to", "xlsx:Calc MS Excel 2007 XML", "--outdir", temporary, str(path.resolve())],
            capture_output=True,
            text=True,
            timeout=max(1, int(remaining)),
        )
        converted = Path(temporary) / path.name
        if completed.returncode != 0 or not converted.exists():
            return False
        shutil.copy2(converted, path)
    return True


def _spreadsheet_case(
    client: OpenAICompatibleClient,
    task: EvalTask,
    input_path: Path,
    output_path: Path,
    config: RunConfig,
    deadline: float,
) -> tuple[list[ApiResponse], list[dict[str, Any]]]:
    shutil.copy2(input_path, output_path)
    messages: list[dict[str, Any]] = [
        {
            "role": "system",
            "content": "You are a spreadsheet agent. Inspect and edit the workbook only with spreadsheet_action. Preserve unrelated cells and call finish when done.",
        },
        {
            "role": "user",
            "content": f"Instruction:\n{task.prompt}\n\nTarget answer position: {task.grader_payload['answer_position']}\n\n{_workbook_summary(output_path)}",
        },
    ]
    responses = []
    actions = []
    calls_used = 0
    while calls_used < config.max_tool_calls:
        _remaining_timeout(client, deadline)
        response = client.complete(messages, [SPREADSHEET_TOOL], config.max_output_tokens)
        responses.append(response)
        messages.append({"role": "assistant", "content": response.text or None, "tool_calls": response.raw_message.get("tool_calls") or []})
        if not response.tool_calls:
            break
        finished = False
        for call in response.tool_calls:
            calls_used += 1
            try:
                arguments = json.loads(call["arguments"]) if isinstance(call["arguments"], str) else call["arguments"]
                result = _spreadsheet_action(output_path, arguments)
            except Exception as exc:
                arguments = {}
                result = {"error": str(exc)}
            actions.append({"arguments": arguments, "result": result})
            messages.append({"role": "tool", "tool_call_id": call.get("id") or f"call-{calls_used}", "content": json.dumps(result, default=str)[: config.max_tool_output_chars]})
            finished = finished or bool(result.get("finished"))
            if calls_used >= config.max_tool_calls:
                break
        if finished:
            break
    return responses, actions


def _gaia_case(
    client: OpenAICompatibleClient,
    task: EvalTask,
    config: RunConfig,
    deadline: float,
) -> tuple[list[ApiResponse], list[dict[str, Any]]]:
    messages: list[dict[str, Any]] = [{"role": "user", "content": _asset_content(task)}]
    responses: list[ApiResponse] = []
    actions: list[dict[str, Any]] = []
    calls_used = 0
    while calls_used < config.max_tool_calls:
        timeout = _remaining_timeout(client, deadline)
        response = client.complete(messages, GAIA_TOOLS, config.max_output_tokens)
        responses.append(response)
        messages.append({"role": "assistant", "content": response.text or None, "tool_calls": response.raw_message.get("tool_calls") or []})
        if not response.tool_calls:
            break
        for call in response.tool_calls:
            calls_used += 1
            try:
                arguments = json.loads(call["arguments"]) if isinstance(call["arguments"], str) else call["arguments"]
                if call["name"] == "web_search":
                    result = _web_search(str(arguments.get("query", "")), timeout, config.max_tool_output_chars)
                elif call["name"] == "web_fetch":
                    result = _web_fetch(str(arguments.get("url", "")), timeout, config.max_tool_output_chars)
                else:
                    raise ValueError(f"unsupported tool: {call['name']}")
                content = json.dumps({"result": result})
            except Exception as exc:
                arguments = {}
                content = json.dumps({"error": f"{type(exc).__name__}: {exc}"})
            actions.append({"name": call["name"], "arguments": arguments, "result": content[: config.max_tool_output_chars]})
            messages.append({"role": "tool", "tool_call_id": call.get("id") or f"call-{calls_used}", "content": content[: config.max_tool_output_chars]})
            if calls_used >= config.max_tool_calls:
                break
    return responses, actions


def _cost(model: ModelConfig, responses: list[ApiResponse]) -> tuple[int, int, int, float, float]:
    input_tokens = sum(item.input_tokens for item in responses)
    output_tokens = sum(item.output_tokens for item in responses)
    cached = sum(item.cached_input_tokens for item in responses)
    billable_input = max(0, input_tokens - cached)
    cost = billable_input * model.input_per_million_usd / 1_000_000
    cached_price = model.cached_input_per_million_usd
    cost += cached * (cached_price if cached_price is not None else model.input_per_million_usd) / 1_000_000
    cost += output_tokens * model.output_per_million_usd / 1_000_000
    latency = sum(item.latency_ms for item in responses)
    return input_tokens, output_tokens, cached, cost, latency


def run_one(
    run_id: str,
    task: EvalTask,
    model: ModelConfig,
    trial: int,
    config: RunConfig,
    run_dir: Path,
) -> EvalResult:
    client = OpenAICompatibleClient(model, config.timeout_seconds)
    artifact_dir = run_dir / "artifacts" / task.eval_id.replace(":", "__") / model.alias / f"trial-{trial}"
    artifact_dir.mkdir(parents=True, exist_ok=True)
    responses: list[ApiResponse] = []
    grade = GradeResult(0.0, 0.0, False)
    error: str | None = None
    artifacts: list[str] = []
    trace: dict[str, Any] = {"task": task.eval_id, "model": model.alias, "trial": trial}
    deadline = time.monotonic() + config.timeout_seconds
    try:
        if task.source == "bfcl":
            response = client.complete(task.metadata["messages"], _openai_tools(task.metadata["functions"]), config.max_output_tokens)
            responses.append(response)
            grade = grade_bfcl(task, response.tool_calls)
            trace["tool_calls"] = response.tool_calls
            trace["response_text"] = response.text
        elif task.source == "gaia":
            responses, actions = _gaia_case(client, task, config, deadline)
            response_text = responses[-1].text if responses else ""
            grade = grade_gaia(task, response_text)
            trace["response_text"] = response_text
            trace["actions"] = actions
        elif task.source == "spreadsheetbench":
            output_files = []
            case_actions = []
            with tempfile.TemporaryDirectory(prefix="task-router-eval-") as temporary:
                temp = Path(temporary)
                for index, asset in enumerate(task.assets, start=1):
                    working = temp / f"case-{index}.xlsx"
                    case_responses, actions = _spreadsheet_case(client, task, Path(asset), working, config, deadline)
                    actions.append({"workbook_recalculated": _recalculate_workbook(working, deadline)})
                    responses.extend(case_responses)
                    case_actions.append(actions)
                    saved = artifact_dir / f"case-{index}.xlsx"
                    shutil.copy2(working, saved)
                    output_files.append(saved)
                    artifacts.append(str(saved.relative_to(run_dir)))
            grade = grade_spreadsheet(task, output_files)
            trace["actions"] = case_actions
        else:
            raise ValueError(f"Unsupported eval source: {task.source}")
    except Exception as exc:
        error = f"{type(exc).__name__}: {exc}"
        grade = GradeResult(0.0, 0.0, False, severe_failure=False, details={"execution_error": error})
    trace_path = artifact_dir / "trace.json"
    trace_path.write_text(json.dumps(trace, indent=2, default=str), encoding="utf-8")
    artifacts.append(str(trace_path.relative_to(run_dir)))
    input_tokens, output_tokens, cached, cost, latency = _cost(model, responses)
    return EvalResult(
        run_id=run_id,
        eval_id=task.eval_id,
        source=task.source,
        category=task.category,
        split=task.split,
        model_alias=model.alias,
        catalog_model_id=model.catalog_model_id,
        trial=trial,
        native_score=grade.native_score,
        normalized_score=grade.normalized_score,
        passed=grade.passed,
        severe_failure=grade.severe_failure,
        latency_ms=latency,
        input_tokens=input_tokens,
        output_tokens=output_tokens,
        cached_input_tokens=cached,
        cost_usd=cost,
        artifact_paths=artifacts,
        error=error,
        grader_details=grade.details,
        provenance={
            "source_revision": task.source_revision,
            "grader": task.grader,
            "api_model": model.api_model,
            "capability_tags": task.capability_tags,
            "temperature": 0,
            "seed": 42,
        },
    )


def _append_result(path: Path, result: EvalResult) -> None:
    with path.open("a", encoding="utf-8") as handle:
        handle.write(json.dumps(result.to_dict(), sort_keys=True) + "\n")


def run_profile(
    profile: str,
    models_config_path: Path,
    cache_dir: Path = DEFAULT_CACHE_DIR,
    artifact_root: Path = ARTIFACT_DIR / "evals",
) -> Path:
    config, models = load_models_config(models_config_path)
    tasks = load_built_tasks(profile, cache_dir)
    build_metadata = json.loads((cache_dir / "built" / profile / "build.json").read_text(encoding="utf-8"))
    started = datetime.now(timezone.utc)
    fingerprint = hashlib.sha256(
        (profile + json.dumps([asdict(model) for model in models], sort_keys=True) + started.isoformat()).encode()
    ).hexdigest()[:8]
    run_id = started.strftime("%Y%m%dT%H%M%SZ") + f"-{fingerprint}"
    run_dir = artifact_root / run_id
    if run_dir.exists():
        raise FileExistsError(f"Eval run already exists: {run_id}")
    run_dir.mkdir(parents=True)
    results_path = run_dir / "results.jsonl"
    metadata = {
        "run_id": run_id,
        "profile": profile,
        "status": "running",
        "started_at": started.isoformat(),
        "task_count": len(tasks),
        "source_counts": dict(sorted(Counter(task.source for task in tasks).items())),
        "build": build_metadata,
        "scaffold_version": "task-router-evals-v1",
        "models": [{key: value for key, value in asdict(model).items() if key != "api_key_env"} for model in models],
        "run_config": asdict(config),
        "experimental": True,
    }
    (run_dir / "run.json").write_text(json.dumps(metadata, indent=2), encoding="utf-8")
    initial: list[EvalResult] = []
    try:
        for task in tasks:
            for model in models:
                result = run_one(run_id, task, model, 0, config, run_dir)
                initial.append(result)
                _append_result(results_path, result)
        by_task: dict[str, list[EvalResult]] = {}
        for result in initial:
            by_task.setdefault(result.eval_id, []).append(result)
        repeat_ids = set()
        for task in tasks:
            values = by_task[task.eval_id]
            disagreement = len({result.passed for result in values}) > 1
            has_error = any(result.error for result in values)
            audit = int(hashlib.sha256(f"42:audit:{task.eval_id}".encode()).hexdigest(), 16) % 10 == 0
            if disagreement or has_error or audit:
                repeat_ids.add(task.eval_id)
        for task in tasks:
            if task.eval_id not in repeat_ids:
                continue
            for trial in (1, 2):
                for model in models:
                    _append_result(results_path, run_one(run_id, task, model, trial, config, run_dir))
        metadata.update(
            {
                "status": "complete",
                "completed_at": datetime.now(timezone.utc).isoformat(),
                "repeated_task_count": len(repeat_ids),
            }
        )
    except BaseException as exc:
        metadata.update({"status": "failed", "completed_at": datetime.now(timezone.utc).isoformat(), "error": str(exc)})
        raise
    finally:
        (run_dir / "run.json").write_text(json.dumps(metadata, indent=2), encoding="utf-8")
    return run_dir
