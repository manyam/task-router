from __future__ import annotations

import datetime as dt
import json
import math
import re
import unicodedata
from pathlib import Path
from typing import Any

from .schema import EvalTask, GradeResult


def _standard_string(value: str) -> str:
    return re.sub(r"[ ,./\-_*^]", "", value.lower().replace("'", '"'))


def _value_matches(actual: Any, options: list[Any]) -> bool:
    if "" in options and actual in {None, ""}:
        return True
    for expected in options:
        if expected == "":
            continue
        if isinstance(expected, str) and isinstance(actual, str):
            if _standard_string(expected) == _standard_string(actual):
                return True
        elif isinstance(expected, float) and isinstance(actual, int):
            if expected == float(actual):
                return True
        elif isinstance(expected, list) and isinstance(actual, (list, tuple)):
            if len(expected) == len(actual) and all(
                _value_matches(item, [target]) for item, target in zip(actual, expected)
            ):
                return True
        elif isinstance(expected, dict) and isinstance(actual, dict):
            if set(actual) == set(expected) and all(
                _value_matches(actual[key], [expected[key]]) for key in actual
            ):
                return True
        elif actual == expected:
            return True
    return False


def _normalize_call(call: dict[str, Any]) -> tuple[str, dict[str, Any]]:
    name = str(call.get("name") or call.get("function", {}).get("name") or "")
    arguments = call.get("arguments", call.get("function", {}).get("arguments", {}))
    if isinstance(arguments, str):
        try:
            arguments = json.loads(arguments)
        except json.JSONDecodeError:
            arguments = {}
    return name, arguments if isinstance(arguments, dict) else {}


def _call_matches(call: dict[str, Any], expected: dict[str, Any]) -> tuple[bool, str]:
    actual_name, actual_args = _normalize_call(call)
    expected_name, expected_args = next(iter(expected.items()))
    if actual_name.replace("_", ".") != expected_name and actual_name != expected_name.replace(".", "_"):
        return False, f"expected function {expected_name}, got {actual_name or '<none>'}"
    for key, options in expected_args.items():
        if key not in actual_args:
            if "" in options:
                continue
            return False, f"missing parameter {key}"
        if not _value_matches(actual_args[key], options):
            return False, f"invalid value for parameter {key}"
    unexpected = set(actual_args) - set(expected_args)
    if unexpected:
        return False, f"unexpected parameters: {sorted(unexpected)}"
    return True, ""


def grade_bfcl(task: EvalTask, tool_calls: list[dict[str, Any]]) -> GradeResult:
    expected = task.grader_payload["expected_calls"]
    if task.category == "irrelevance":
        passed = len(tool_calls) == 0
        return GradeResult(
            native_score=float(passed),
            normalized_score=float(passed),
            passed=passed,
            details={"expected_call_count": 0, "actual_call_count": len(tool_calls)},
        )
    if len(tool_calls) != len(expected):
        return GradeResult(
            native_score=0.0,
            normalized_score=0.0,
            passed=False,
            details={
                "error": "wrong number of function calls",
                "expected_call_count": len(expected),
                "actual_call_count": len(tool_calls),
            },
        )
    unused = set(range(len(tool_calls)))
    errors: list[str] = []
    for expected_call in expected:
        match = None
        candidate_errors = []
        for index in sorted(unused):
            valid, error = _call_matches(tool_calls[index], expected_call)
            if valid:
                match = index
                break
            candidate_errors.append(error)
        if match is None:
            errors.append(candidate_errors[0] if candidate_errors else "missing call")
        else:
            unused.remove(match)
    passed = not errors and not unused
    return GradeResult(
        native_score=float(passed),
        normalized_score=float(passed),
        passed=passed,
        details={"errors": errors, "unmatched_call_indices": sorted(unused)},
    )


def _excel_value(value: Any) -> Any:
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return round(float(value), 2)
    if isinstance(value, dt.time):
        return value.strftime("%H:%M")
    if isinstance(value, dt.datetime):
        if value.tzinfo is not None:
            value = value.astimezone(dt.timezone.utc).replace(tzinfo=None)
        start = dt.datetime(1899, 12, 30)
        return round((value - start).total_seconds() / 86400.0, 0)
    if isinstance(value, dt.date):
        start = dt.date(1899, 12, 30)
        return float((value - start).days)
    if isinstance(value, str):
        try:
            return round(float(value), 2)
        except ValueError:
            return value
    if value == "":
        return None
    return value


def _cell_names(cell_range: str) -> list[str]:
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    if ":" not in cell_range:
        return [cell_range]
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    return [
        f"{get_column_letter(column)}{row}"
        for column in range(min_col, max_col + 1)
        for row in range(min_row, max_row + 1)
    ]


def compare_workbooks(
    answer_file: Path,
    output_file: Path,
    answer_position: str,
    answer_sheet: str | None = None,
) -> tuple[bool, str]:
    if not output_file.exists():
        return False, "output file is missing"
    expected = None
    actual = None
    try:
        import openpyxl

        expected = openpyxl.load_workbook(answer_file, data_only=True)
        actual = openpyxl.load_workbook(output_file, data_only=True)
    except Exception as exc:
        if expected is not None:
            expected.close()
        return False, f"workbook could not be opened: {exc}"
    try:
        default_sheet = answer_sheet.strip("'") if answer_sheet else expected.sheetnames[0]
        for segment in answer_position.split(","):
            segment = segment.strip()
            if "!" in segment:
                sheet, cell_range = segment.split("!", 1)
                sheet = sheet.strip("'")
            else:
                sheet, cell_range = default_sheet, segment
            cell_range = cell_range.strip("'").replace(" ", "")
            if sheet not in expected.sheetnames or sheet not in actual.sheetnames:
                return False, f"worksheet is missing: {sheet}"
            for cell_name in _cell_names(cell_range):
                left = _excel_value(expected[sheet][cell_name].value)
                right = _excel_value(actual[sheet][cell_name].value)
                if left != right:
                    return False, f"value differs at {sheet}!{cell_name}"
        return True, ""
    finally:
        expected.close()
        actual.close()


def grade_spreadsheet(task: EvalTask, output_files: list[Path]) -> GradeResult:
    answers = [Path(value) for value in task.grader_payload["answer_files"]]
    if len(output_files) != len(answers):
        return GradeResult(0.0, 0.0, False, True, {"error": "wrong output file count"})
    case_results = []
    errors = []
    severe = False
    answer_sheet = task.grader_payload.get("answer_sheet")
    for answer, output in zip(answers, output_files):
        passed, error = compare_workbooks(
            answer,
            output,
            task.grader_payload["answer_position"],
            answer_sheet=answer_sheet,
        )
        case_results.append(int(passed))
        errors.append(error)
        severe = severe or "could not be opened" in error
    score = sum(case_results) / len(case_results)
    return GradeResult(
        native_score=score,
        normalized_score=score,
        passed=score == 1.0,
        severe_failure=severe,
        details={"test_case_results": case_results, "errors": errors},
    )


def _normalize_gaia(value: str) -> str:
    value = unicodedata.normalize("NFKC", value).strip().lower()
    value = re.sub(r"^(final answer|answer)\s*:\s*", "", value)
    value = value.strip().strip("'\"")
    value = re.sub(r"\s+", " ", value)
    value = re.sub(r"\s*,\s*", ",", value)
    return value.rstrip(".")


def grade_gaia(task: EvalTask, response_text: str) -> GradeResult:
    expected = _normalize_gaia(str(task.grader_payload["final_answer"]))
    actual = _normalize_gaia(response_text)
    passed = actual == expected
    if not passed:
        try:
            passed = math.isclose(float(actual.replace(",", "")), float(expected.replace(",", "")), rel_tol=1e-6)
        except ValueError:
            pass
    return GradeResult(
        native_score=float(passed),
        normalized_score=float(passed),
        passed=passed,
        details={"normalized_response": actual},
    )
