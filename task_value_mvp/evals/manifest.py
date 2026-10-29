from __future__ import annotations

import hashlib
import json
import os
import re
import shutil
import tarfile
import tempfile
import urllib.request
import warnings
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Iterable

from .schema import EvalTask


BASE_DIR = Path(__file__).resolve().parents[2]
MANIFEST_PATH = BASE_DIR / "evals" / "manifest.json"
DEFAULT_CACHE_DIR = Path(os.environ.get("TASK_ROUTER_EVAL_CACHE", BASE_DIR / ".eval-cache"))

BFCL_FILES = {
    "simple": "BFCL_v4_simple_python.json",
    "multiple": "BFCL_v4_multiple.json",
    "parallel": "BFCL_v4_parallel.json",
    "irrelevance": "BFCL_v4_irrelevance.json",
}


def load_manifest(path: Path = MANIFEST_PATH) -> dict[str, Any]:
    value = json.loads(path.read_text(encoding="utf-8"))
    if value.get("version") != 1:
        raise ValueError("Unsupported eval manifest version")
    return value


def _selected(source: dict[str, Any]) -> Iterable[tuple[str, str, str]]:
    for stratum, spec in source["strata"].items():
        holdout = set(spec["holdout_ids"])
        for source_id in spec["selected_ids"]:
            yield stratum, str(source_id), "holdout" if source_id in holdout else "calibration"


def _expected_selection(source: dict[str, Any], stratum: str, population: Iterable[str], count: int) -> list[str]:
    key = source["selection_key"]
    return sorted(
        (str(value) for value in population),
        key=lambda source_id: hashlib.sha256(
            key.format(stratum=stratum, source_task_id=source_id).encode()
        ).hexdigest(),
    )[:count]


def _jsonl(path: Path) -> list[dict[str, Any]]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def _fetch(url: str, destination: Path) -> None:
    if destination.exists():
        return
    destination.parent.mkdir(parents=True, exist_ok=True)
    request = urllib.request.Request(url, headers={"User-Agent": "task-router-evals/1"})
    with urllib.request.urlopen(request, timeout=60) as response, destination.open("wb") as handle:
        shutil.copyfileobj(response, handle)


def _verify_file(path: Path, expected: str, label: str) -> None:
    if not path.exists():
        raise ValueError(f"Pinned source asset is missing: {label}")
    actual = hashlib.sha256(path.read_bytes()).hexdigest()
    if actual != expected:
        raise ValueError(f"Pinned source asset changed: {label} (expected {expected}, got {actual})")


def _build_bfcl(source: dict[str, Any], cache_dir: Path) -> list[EvalTask]:
    root = cache_dir / "sources" / "bfcl" / source["revision"]
    base = (
        "https://raw.githubusercontent.com/ShishirPatil/gorilla/"
        f"{source['revision']}/berkeley-function-call-leaderboard/bfcl_eval/data"
    )
    questions: dict[str, dict[str, Any]] = {}
    answers: dict[str, dict[str, Any]] = {}
    population_ids: dict[str, list[str]] = {}
    for category, filename in BFCL_FILES.items():
        question_path = root / filename
        _fetch(f"{base}/{filename}", question_path)
        _verify_file(question_path, source["file_sha256"][filename], f"BFCL {filename}")
        category_questions = _jsonl(question_path)
        population_ids[category] = [item["id"] for item in category_questions]
        questions.update({item["id"]: item for item in category_questions})
        if category != "irrelevance":
            answer_path = root / "possible_answer" / filename
            _fetch(f"{base}/possible_answer/{filename}", answer_path)
            relative_answer = f"possible_answer/{filename}"
            _verify_file(answer_path, source["file_sha256"][relative_answer], f"BFCL {relative_answer}")
            answers.update({item["id"]: item for item in _jsonl(answer_path)})

    for category, spec in source["strata"].items():
        expected = _expected_selection(source, category, population_ids[category], len(spec["selected_ids"]))
        if expected != spec["selected_ids"]:
            raise ValueError(f"BFCL deterministic selection changed for {category}")

    tasks = []
    for category, source_id, split in _selected(source):
        if source_id not in questions:
            raise ValueError(f"Pinned BFCL task is missing: {source_id}")
        item = questions[source_id]
        messages = item["question"][0]
        prompt = "\n".join(str(message.get("content", "")) for message in messages).strip()
        expected = [] if category == "irrelevance" else answers.get(source_id, {}).get("ground_truth")
        if expected is None:
            raise ValueError(f"Pinned BFCL answer is missing: {source_id}")
        tasks.append(
            EvalTask(
                eval_id=f"bfcl:{source_id}",
                source="bfcl",
                source_revision=source["revision"],
                source_task_id=source_id,
                prompt=prompt,
                category=category,
                capability_tags=["tool_calling", category],
                grader=source["grader"],
                split=split,
                access_policy=source["access_policy"],
                grader_payload={"expected_calls": expected},
                metadata={"messages": messages, "functions": item["function"]},
            )
        )
    return tasks


def _safe_extract(archive: Path, destination: Path) -> None:
    destination.mkdir(parents=True, exist_ok=True)
    resolved = destination.resolve()
    with tarfile.open(archive, "r:gz") as tar:
        for member in tar.getmembers():
            target = (destination / member.name).resolve()
            if resolved not in target.parents and target != resolved:
                raise ValueError(f"Unsafe archive member: {member.name}")
        tar.extractall(destination)


def _hf_download(repo_id: str, filename: str, revision: str, local_dir: Path) -> Path:
    try:
        from huggingface_hub import hf_hub_download
    except ImportError as exc:
        raise RuntimeError("Install requirements.txt before building eval datasets") from exc
    local_dir.mkdir(parents=True, exist_ok=True)
    return Path(
        hf_hub_download(
            repo_id=repo_id,
            filename=filename,
            repo_type="dataset",
            revision=revision,
            local_dir=str(local_dir),
        )
    )


def _build_spreadsheetbench(source: dict[str, Any], cache_dir: Path) -> list[EvalTask]:
    source_dir = cache_dir / "sources" / "spreadsheetbench" / source["revision"]
    archive = _hf_download(source["dataset_id"], source["filename"], source["revision"], source_dir)
    _verify_file(archive, source["file_sha256"][source["filename"]], "SpreadsheetBench archive")
    extracted = source_dir / "extracted"
    dataset_root = extracted / "spreadsheetbench_verified_400"
    if not (dataset_root / "dataset.json").exists():
        _safe_extract(archive, extracted)
    rows = {str(item["id"]): item for item in json.loads((dataset_root / "dataset.json").read_text())}
    if len(rows) != int(source["population_count"]):
        raise ValueError(f"SpreadsheetBench population changed: expected {source['population_count']}, got {len(rows)}")
    _validate_spreadsheet_strata(source, dataset_root, rows)

    tasks = []
    for stratum, source_id, split in _selected(source):
        item = rows.get(source_id)
        if item is None:
            raise ValueError(f"Pinned SpreadsheetBench task is missing: {source_id}")
        task_dir = dataset_root / item["spreadsheet_path"]
        inputs = (
            sorted(task_dir.glob(f"*_{source_id}_input.xlsx"))
            or sorted(task_dir.glob(f"*_{source_id}_init.xlsx"))
            or sorted(task_dir.glob("initial.xlsx"))
        )
        answers = (
            sorted(task_dir.glob(f"*_{source_id}_answer.xlsx"))
            or sorted(task_dir.glob(f"*_{source_id}_golden.xlsx"))
            or sorted(task_dir.glob("*golden.xlsx"))
        )
        if not inputs or len(inputs) != len(answers):
            raise ValueError(f"SpreadsheetBench task {source_id} has incomplete input/answer files")
        tasks.append(
            EvalTask(
                eval_id=f"spreadsheetbench:{source_id}",
                source="spreadsheetbench",
                source_revision=source["revision"],
                source_task_id=source_id,
                prompt=item["instruction"].strip(),
                category=item["instruction_type"],
                capability_tags=["spreadsheet", stratum, item["instruction_type"].lower().replace(" ", "_")],
                grader=source["grader"],
                split=split,
                access_policy=source["access_policy"],
                modalities=["text", "spreadsheet"],
                assets=[str(path.resolve()) for path in inputs],
                grader_payload={
                    "answer_files": [str(path.resolve()) for path in answers],
                    "answer_position": item["answer_position"],
                    "answer_sheet": item.get("answer_sheet"),
                    "instruction_type": item["instruction_type"],
                },
                metadata={"stratum": stratum, "answer_sheet": item.get("answer_sheet")},
            )
        )
    return tasks


def _validate_spreadsheet_strata(
    source: dict[str, Any],
    dataset_root: Path,
    rows: dict[str, dict[str, Any]],
) -> None:
    import openpyxl

    populations: dict[str, list[tuple[str, int]]] = {"cell": [], "sheet": []}
    for source_id, item in rows.items():
        task_dir = dataset_root / item["spreadsheet_path"]
        inputs = (
            sorted(task_dir.glob(f"*_{source_id}_input.xlsx"))
            or sorted(task_dir.glob(f"*_{source_id}_init.xlsx"))
            or sorted(task_dir.glob("initial.xlsx"))
        )
        answers = (
            sorted(task_dir.glob(f"*_{source_id}_answer.xlsx"))
            or sorted(task_dir.glob(f"*_{source_id}_golden.xlsx"))
            or sorted(task_dir.glob("*golden.xlsx"))
        )
        if not inputs or len(inputs) != len(answers):
            raise ValueError(f"SpreadsheetBench population task has incomplete assets: {source_id}")
        with warnings.catch_warnings():
            warnings.simplefilter("ignore", UserWarning)
            workbook = openpyxl.load_workbook(inputs[0], read_only=True, data_only=False)
            nonempty = 0
            formulas = 0
            for worksheet in workbook.worksheets:
                for row in worksheet.iter_rows():
                    for cell in row:
                        if cell.value is not None:
                            nonempty += 1
                            formulas += int(isinstance(cell.value, str) and cell.value.startswith("="))
        score = nonempty + 5 * formulas + 25 * len(workbook.sheetnames)
        workbook.close()
        task_type = "cell" if item["instruction_type"].startswith("Cell") else "sheet"
        populations[task_type].append((source_id, score))

    assigned: dict[str, str] = {}
    stratum_members: dict[str, list[str]] = {}
    for task_type, population in populations.items():
        ordered = sorted(population, key=lambda value: (value[1], value[0]))
        first, second = len(ordered) // 3, 2 * len(ordered) // 3
        for tier, values in zip(("low", "medium", "high"), (ordered[:first], ordered[first:second], ordered[second:])):
            stratum = f"{task_type}_{tier}"
            if len(values) != int(source["strata"][stratum]["population_count"]):
                raise ValueError(f"SpreadsheetBench {stratum} population changed")
            stratum_members[stratum] = [source_id for source_id, _ in values]
            assigned.update({source_id: stratum for source_id, _ in values})
    for stratum, spec in source["strata"].items():
        wrong = [source_id for source_id in spec["selected_ids"] if assigned.get(source_id) != stratum]
        if wrong:
            raise ValueError(f"SpreadsheetBench tasks left their complexity stratum {stratum}: {wrong}")
        expected = _expected_selection(source, stratum, stratum_members[stratum], len(spec["selected_ids"]))
        if expected != spec["selected_ids"]:
            raise ValueError(f"SpreadsheetBench deterministic selection changed for {stratum}")


def _build_gaia(source: dict[str, Any], cache_dir: Path) -> list[EvalTask]:
    source_dir = cache_dir / "sources" / "gaia" / source["revision"]
    metadata_path = _hf_download(
        source["dataset_id"], "2023/validation/metadata.parquet", source["revision"], source_dir
    )
    _verify_file(
        metadata_path,
        source["file_sha256"]["2023/validation/metadata.parquet"],
        "GAIA validation metadata",
    )
    try:
        import duckdb
    except ImportError as exc:
        raise RuntimeError("duckdb is required to read GAIA metadata") from exc
    selected = {source_id for _, source_id, _ in _selected(source)}
    placeholders = ",".join("?" for _ in selected)
    population_rows = duckdb.connect().execute(
        'SELECT task_id, Level, file_name FROM read_parquet(?)', [str(metadata_path)]
    ).fetchall()
    supported_extensions = tuple(source["supported_file_extensions"])
    for level, spec in source["strata"].items():
        population = [
            str(task_id)
            for task_id, task_level, file_name in population_rows
            if str(task_level) == level
            and (not file_name or Path(str(file_name)).suffix.lower() in supported_extensions)
        ]
        if len(population) != int(spec["population_count"]):
            raise ValueError(f"GAIA supported population changed for level {level}")
        expected = _expected_selection(source, level, population, len(spec["selected_ids"]))
        if expected != spec["selected_ids"]:
            raise ValueError(f"GAIA deterministic selection changed for level {level}")
    query = f'''SELECT task_id, Question, Level, "Final answer", file_name
                FROM read_parquet(?) WHERE task_id IN ({placeholders})'''
    rows = duckdb.connect().execute(query, [str(metadata_path), *sorted(selected)]).fetchall()
    by_id = {row[0]: row for row in rows}
    tasks = []
    for level, source_id, split in _selected(source):
        if source_id not in by_id:
            raise ValueError(f"Pinned GAIA task is missing: {source_id}")
        _, question, _, final_answer, file_name = by_id[source_id]
        assets: list[str] = []
        modalities = ["text"]
        if file_name:
            path = _hf_download(
                source["dataset_id"], f"2023/validation/{file_name}", source["revision"], source_dir
            )
            relative_asset = f"2023/validation/{file_name}"
            _verify_file(path, source["file_sha256"][relative_asset], f"GAIA {relative_asset}")
            assets.append(str(path.resolve()))
            suffix = path.suffix.lower()
            modalities.append("vision" if suffix in {".png", ".jpg", ".jpeg"} else "file")
        tasks.append(
            EvalTask(
                eval_id=f"gaia:{source_id}",
                source="gaia",
                source_revision=source["revision"],
                source_task_id=source_id,
                prompt=question.strip(),
                category=f"level_{level}",
                capability_tags=["general_assistant", "research", f"level_{level}"],
                grader=source["grader"],
                split=split,
                access_policy=source["access_policy"],
                modalities=modalities,
                assets=assets,
                grader_payload={"final_answer": final_answer},
                metadata={"level": int(level), "file_name": file_name or None},
            )
        )
    return tasks


def _normalize_prompt(value: str) -> str:
    return " ".join(re.findall(r"[a-z0-9]+", value.lower()))


def _ngrams(value: str, size: int = 5) -> set[tuple[str, ...]]:
    words = value.split()
    if len(words) < size:
        return {tuple(words)}
    return {tuple(words[index : index + size]) for index in range(len(words) - size + 1)}


def validate_tasks(tasks: list[EvalTask], threshold: float) -> None:
    ids = [task.eval_id for task in tasks]
    duplicates = [key for key, count in Counter(ids).items() if count > 1]
    if duplicates:
        raise ValueError(f"Duplicate eval IDs: {duplicates}")
    normalized = [_normalize_prompt(task.prompt) for task in tasks]
    if len(normalized) != len(set(normalized)):
        raise ValueError("Exact duplicate prompts found in eval manifest")
    grams = [_ngrams(value) for value in normalized]
    for left in range(len(tasks)):
        for right in range(left + 1, len(tasks)):
            union = grams[left] | grams[right]
            similarity = len(grams[left] & grams[right]) / max(1, len(union))
            if similarity >= threshold:
                raise ValueError(
                    f"Near-duplicate prompts ({similarity:.3f}): {tasks[left].eval_id}, {tasks[right].eval_id}"
                )


def build_profile(profile: str, cache_dir: Path = DEFAULT_CACHE_DIR) -> tuple[Path, dict[str, Any]]:
    manifest = load_manifest()
    if profile not in manifest["profiles"]:
        raise ValueError(f"Unknown eval profile: {profile}")
    builders = {"bfcl": _build_bfcl, "spreadsheetbench": _build_spreadsheetbench, "gaia": _build_gaia}
    tasks: list[EvalTask] = []
    for source_name in manifest["profiles"][profile]:
        source = manifest["sources"][source_name]
        source_tasks = builders[source_name](source, cache_dir)
        expected = sum(len(spec["selected_ids"]) for spec in source["strata"].values())
        if len(source_tasks) != expected:
            raise ValueError(f"{source_name} returned {len(source_tasks)} tasks; expected {expected}")
        tasks.extend(source_tasks)
    validate_tasks(tasks, float(manifest["near_duplicate_threshold"]))

    output_dir = cache_dir / "built" / profile
    output_dir.mkdir(parents=True, exist_ok=True)
    task_path = output_dir / "tasks.jsonl"
    with tempfile.NamedTemporaryFile("w", encoding="utf-8", dir=output_dir, delete=False) as handle:
        temp_path = Path(handle.name)
        for task in tasks:
            handle.write(json.dumps(task.to_dict(), sort_keys=True) + "\n")
    temp_path.replace(task_path)
    counts = Counter(task.source for task in tasks)
    split_counts = Counter(task.split for task in tasks)
    build = {
        "profile": profile,
        "built_at": datetime.now(timezone.utc).isoformat(),
        "task_count": len(tasks),
        "source_counts": dict(sorted(counts.items())),
        "split_counts": dict(sorted(split_counts.items())),
        "manifest_sha256": hashlib.sha256(MANIFEST_PATH.read_bytes()).hexdigest(),
        "source_revisions": {
            name: manifest["sources"][name]["revision"] for name in manifest["profiles"][profile]
        },
    }
    (output_dir / "build.json").write_text(json.dumps(build, indent=2), encoding="utf-8")
    return task_path, build


def load_built_tasks(profile: str, cache_dir: Path = DEFAULT_CACHE_DIR) -> list[EvalTask]:
    path = cache_dir / "built" / profile / "tasks.jsonl"
    build_path = path.parent / "build.json"
    if not path.exists() or not build_path.exists():
        raise FileNotFoundError(f"Eval profile is not built: {profile}. Run the eval build command first.")
    build = json.loads(build_path.read_text(encoding="utf-8"))
    current_manifest = hashlib.sha256(MANIFEST_PATH.read_bytes()).hexdigest()
    if build.get("manifest_sha256") != current_manifest:
        raise ValueError(f"Eval profile {profile} was built from a different manifest; rebuild it before running")
    return [EvalTask.from_dict(value) for value in _jsonl(path)]
